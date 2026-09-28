# export/excel_utils.py

"""توابع کمکی Excel - تنظیمات صفحه، ستون‌ها، و ..."""

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment


class ExcelUtils:
    """کلاس توابع کمکی Excel"""

    def __init__(self, styles=None):
        self.styles = styles

    def set_page_setup_a3_landscape(self, ws):
        """تنظیمات صفحه A3 Landscape"""
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(
            left=0.5,
            right=0.5,
            top=0.75,
            bottom=0.75,
            header=0.3,
            footer=0.3
        )

    def set_repeat_header_rows(self, ws, start_row: int = 1, end_row: int = 2):
        """تنظیم تکرار هدر در همه صفحات"""
        ws.print_title_rows = f'{start_row}:{end_row}'

    def set_column_widths(self, ws, widths: Dict[str, int]):
        """تنظیم عرض ستون‌ها"""
        for col_letter, width in widths.items():
            ws.column_dimensions[col_letter].width = width

    def set_row_height(self, ws, row: int, height: int = 20):
        """تنظیم ارتفاع ردیف"""
        ws.row_dimensions[row].height = height

    def add_auto_filter(self, ws, header_row: int, col_count: int):
        """افزودن AutoFilter"""
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(col_count)}{ws.max_row}"

    def freeze_panes(self, ws, cell: str = "A3"):
        """فریز کردن پنل"""
        ws.freeze_panes = cell

    def add_conditional_formatting_io(self, ws, start_row: int, start_col: int, end_row: int, end_col: int):
        """افزودن قالب‌بندی شرطی برای ستون‌های I/O"""
        start_letter = get_column_letter(start_col)
        end_letter = get_column_letter(end_col)
        cell_range = f"{start_letter}{start_row}:{end_letter}{end_row}"
        
        ws.conditional_formatting.add(
            cell_range,
            ColorScaleRule(
                start_type='min',
                start_color='F8D7DA',
                mid_type='percentile',
                mid_value=50,
                mid_color='FFF3CD',
                end_type='max',
                end_color='D4EDDA'
            )
        )

    def add_conditional_formatting_status(self, ws, start_row: int, start_col: int, end_row: int, end_col: int):
        """افزودن قالب‌بندی شرطی برای ستون‌های Status"""
        start_letter = get_column_letter(start_col)
        end_letter = get_column_letter(end_col)
        cell_range = f"{start_letter}{start_row}:{end_letter}{end_row}"
        
        ws.conditional_formatting.add(
            cell_range,
            ColorScaleRule(
                start_type='min',
                start_color='FFFFFF',
                end_type='max',
                end_color='1F4E79'
            )
        )

    def merge_and_center(self, ws, start_row: int, start_col: int, end_row: int, end_col: int, value: str = None):
        """ادغام و وسط‌چین کردن سلول‌ها"""
        ws.merge_cells(
            start_row=start_row,
            start_column=start_col,
            end_row=end_row,
            end_column=end_col
        )
        if value:
            cell = ws.cell(row=start_row, column=start_col)
            cell.value = value

    def set_border(self, ws, start_row: int, start_col: int, end_row: int, end_col: int, border):
        """تنظیم مرز برای محدوده سلول‌ها"""
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).border = border

    def set_alignment(self, ws, start_row: int, start_col: int, end_row: int, end_col: int, alignment):
        """تنظیم ترازبندی برای محدوده سلول‌ها"""
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).alignment = alignment

    def set_fill(self, ws, start_row: int, start_col: int, end_row: int, end_col: int, fill):
        """تنظیم پس‌زمینه برای محدوده سلول‌ها"""
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                ws.cell(row=row, column=col).fill = fill

    def write_header_row(self, ws, row: int, headers: list, styles=None):
        """نوشتن ردیف هدر با استایل"""
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            if styles:
                styles.apply_header_style(ws, row, len(headers))

    def write_data_row(self, ws, row: int, data: list, styles=None, is_alt: bool = False):
        """نوشتن ردیف داده با استایل"""
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = value
        if styles:
            styles.apply_body_style(ws, row, len(data), is_alt=is_alt)

    def write_total_row(self, ws, row: int, data: list, styles=None):
        """نوشتن ردیف TOTAL با استایل"""
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = value
        if styles:
            styles.apply_total_style(ws, row, len(data))

    def auto_adjust_column_widths(self, ws, max_width: int = 50):
        """تنظیم خودکار عرض ستون‌ها بر اساس محتوا"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            
            adjusted_width = min(max_length + 2, max_width)
            ws.column_dimensions[column_letter].width = adjusted_width

    def hide_columns(self, ws, columns: list):
        """مخفی کردن ستون‌ها"""
        for col in columns:
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].hidden = True

    def show_columns(self, ws, columns: list):
        """نمایش ستون‌ها"""
        for col in columns:
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].hidden = False

    def add_page_break(self, ws, row: int):
        """افزودن شکست صفحه"""
        ws.row_breaks.append(row)

    def set_zoom(self, ws, zoom: int = 80):
        """تنظیم بزرگنمایی صفحه"""
        ws.sheet_view.zoomScale = zoom

    def set_print_area(self, ws, last_col: str, last_row: int):
        """تنظیم ناحیه چاپ"""
        ws.print_area = f"A1:{last_col}{last_row}"

    def set_footer(self, ws, text: str):
        """تنظیم فوتر صفحه"""
        ws.oddFooter.left.text = text
        ws.oddFooter.left.size = 8
        ws.oddFooter.left.color = "1F4E79"

    def set_header(self, ws, text: str):
        """تنظیم هدر صفحه"""
        ws.oddHeader.center.text = text
        ws.oddHeader.center.size = 10
        ws.oddHeader.center.color = "1F4E79"

    def get_max_row_with_data(self, ws) -> int:
        """دریافت آخرین ردیف دارای داده"""
        return ws.max_row

    def get_max_col_with_data(self, ws) -> int:
        """دریافت آخرین ستون دارای داده"""
        return ws.max_column