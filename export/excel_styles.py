# export/excel_styles.py

"""استایل‌های Excel - تنظیمات ظاهری حرفه‌ای"""

from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle
)
from openpyxl.utils import get_column_letter


class ExcelStyles:
    """کلاس مدیریت استایل‌های Excel"""

    # ===== رنگ‌ها =====
    PRIMARY_COLOR = "1F4E79"
    SECONDARY_COLOR = "2C3E50"
    ACCENT_COLOR = "3498DB"
    SUCCESS_COLOR = "27AE60"
    WARNING_COLOR = "F39C12"
    DANGER_COLOR = "E74C3C"
    LIGHT_COLOR = "F8F9FA"
    WHITE_COLOR = "FFFFFF"
    DARK_COLOR = "2C3E50"
    TOTAL_BG_COLOR = "E6E6FA"
    ALT_ROW_COLOR = "F0F4F8"

    # ===== رنگ‌های ملایم برای ستون‌های Active (مناسب پرینت) =====
    ACTIVE_HEADER_COLOR = "D6EAF8"      # آبی خیلی روشن
    ACTIVE_POSITIVE_COLOR = "E8F5E9"    # سبز خیلی ملایم
    ACTIVE_ZERO_COLOR = "F5F5F5"        # خاکستری خیلی ملایم
    ACTIVE_TOTAL_COLOR = "F3E5F5"       # بنفش خیلی ملایم

    def __init__(self):
        self._create_styles()

    def _create_styles(self):
        """ایجاد تمام استایل‌ها"""

        # ===== فونت‌ها =====
        self.title_font = Font(name="Calibri", size=18, bold=True, color=self.WHITE_COLOR)
        self.subtitle_font = Font(name="Calibri", size=14, bold=True, color=self.PRIMARY_COLOR)
        self.section_font = Font(name="Calibri", size=12, bold=True, color=self.WHITE_COLOR)
        self.subsection_font = Font(name="Calibri", size=11, bold=True, color=self.PRIMARY_COLOR)
        self.header_font = Font(name="Calibri", size=10, bold=True, color=self.WHITE_COLOR)
        self.body_font = Font(name="Calibri", size=9, color="000000")
        self.body_bold_font = Font(name="Calibri", size=9, bold=True, color="000000")
        self.total_font = Font(name="Calibri", size=10, bold=True, color="000000")
        self.small_font = Font(name="Calibri", size=8, color="666666")

        # ===== پس‌زمینه‌ها =====
        self.title_fill = PatternFill(start_color=self.PRIMARY_COLOR, end_color=self.PRIMARY_COLOR, fill_type="solid")
        self.subtitle_fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
        self.section_fill = PatternFill(start_color=self.SECONDARY_COLOR, end_color=self.SECONDARY_COLOR, fill_type="solid")
        self.subsection_fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
        self.header_fill = PatternFill(start_color=self.PRIMARY_COLOR, end_color=self.PRIMARY_COLOR, fill_type="solid")
        self.body_fill = PatternFill(start_color=self.WHITE_COLOR, end_color=self.WHITE_COLOR, fill_type="solid")
        self.alt_fill = PatternFill(start_color=self.ALT_ROW_COLOR, end_color=self.ALT_ROW_COLOR, fill_type="solid")
        self.total_fill = PatternFill(start_color=self.TOTAL_BG_COLOR, end_color=self.TOTAL_BG_COLOR, fill_type="solid")
        self.warning_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        self.danger_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")

        # ===== 🆕 رنگ‌های ملایم برای ستون‌های Active =====
        self.active_header_fill = PatternFill(start_color=self.ACTIVE_HEADER_COLOR, end_color=self.ACTIVE_HEADER_COLOR, fill_type="solid")
        self.active_positive_fill = PatternFill(start_color=self.ACTIVE_POSITIVE_COLOR, end_color=self.ACTIVE_POSITIVE_COLOR, fill_type="solid")
        self.active_zero_fill = PatternFill(start_color=self.ACTIVE_ZERO_COLOR, end_color=self.ACTIVE_ZERO_COLOR, fill_type="solid")
        self.active_total_fill = PatternFill(start_color=self.ACTIVE_TOTAL_COLOR, end_color=self.ACTIVE_TOTAL_COLOR, fill_type="solid")

        # ===== ترازبندی =====
        self.center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        self.right_alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        self.top_left_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        self.top_center_alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

        # ===== مرزها =====
        self.thin_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thin", color="D0D0D0")
        )

        self.medium_border = Border(
            left=Side(style="medium", color=self.PRIMARY_COLOR),
            right=Side(style="medium", color=self.PRIMARY_COLOR),
            top=Side(style="medium", color=self.PRIMARY_COLOR),
            bottom=Side(style="medium", color=self.PRIMARY_COLOR)
        )

        self.thick_bottom_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thick", color=self.PRIMARY_COLOR)
        )

        # ===== استایل‌های نام‌گذاری شده =====
        self.named_styles = {
            'Title': NamedStyle(
                name='Title',
                font=self.title_font,
                fill=self.title_fill,
                alignment=self.center_alignment,
                border=self.medium_border
            ),
            'Subtitle': NamedStyle(
                name='Subtitle',
                font=self.subtitle_font,
                fill=self.subtitle_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'SectionHeader': NamedStyle(
                name='SectionHeader',
                font=self.section_font,
                fill=self.section_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'TableHeader': NamedStyle(
                name='TableHeader',
                font=self.header_font,
                fill=self.header_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'TableBody': NamedStyle(
                name='TableBody',
                font=self.body_font,
                fill=self.body_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'TableBodyAlt': NamedStyle(
                name='TableBodyAlt',
                font=self.body_font,
                fill=self.alt_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'TableBodyLeft': NamedStyle(
                name='TableBodyLeft',
                font=self.body_font,
                fill=self.body_fill,
                alignment=self.left_alignment,
                border=self.thin_border
            ),
            'Total': NamedStyle(
                name='Total',
                font=self.total_font,
                fill=self.total_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'Warning': NamedStyle(
                name='Warning',
                font=self.body_bold_font,
                fill=self.warning_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'Danger': NamedStyle(
                name='Danger',
                font=self.body_bold_font,
                fill=self.danger_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            # ===== 🆕 استایل‌های Active Components =====
            'ActiveHeader': NamedStyle(
                name='ActiveHeader',
                font=Font(name="Calibri", size=9, bold=True, color=self.PRIMARY_COLOR),
                fill=self.active_header_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'ActivePositive': NamedStyle(
                name='ActivePositive',
                font=self.body_font,
                fill=self.active_positive_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'ActiveZero': NamedStyle(
                name='ActiveZero',
                font=self.body_font,
                fill=self.active_zero_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            ),
            'ActiveTotal': NamedStyle(
                name='ActiveTotal',
                font=self.total_font,
                fill=self.active_total_fill,
                alignment=self.center_alignment,
                border=self.thin_border
            )
        }

    def get_style(self, style_name: str) -> NamedStyle:
        """دریافت استایل نام‌گذاری شده"""
        return self.named_styles.get(style_name, self.named_styles['TableBody'])

    def apply_style(self, cell, style_name: str):
        """اعمال استایل به سلول"""
        style = self.get_style(style_name)
        cell.font = style.font
        cell.fill = style.fill
        cell.alignment = style.alignment
        cell.border = style.border

    def apply_header_style(self, ws, row: int, col_count: int):
        """اعمال استایل هدر به یک ردیف"""
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            self.apply_style(cell, 'TableHeader')

    def apply_body_style(self, ws, row: int, col_count: int, is_alt: bool = False):
        """اعمال استایل بدنه به یک ردیف"""
        style_name = 'TableBodyAlt' if is_alt else 'TableBody'
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            self.apply_style(cell, style_name)

    def apply_total_style(self, ws, row: int, col_count: int):
        """اعمال استایل TOTAL به یک ردیف"""
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col)
            self.apply_style(cell, 'Total')

    def apply_active_header_style(self, ws, row: int, start_col: int, end_col: int):
        """اعمال استایل هدر Active Components"""
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            self.apply_style(cell, 'ActiveHeader')

    def apply_active_body_style(self, ws, row: int, start_col: int, end_col: int, values: list):
        """اعمال استایل بدنه Active Components بر اساس مقدار"""
        for i, col in enumerate(range(start_col, end_col + 1)):
            cell = ws.cell(row=row, column=col)
            if values[i] > 0:
                self.apply_style(cell, 'ActivePositive')
            else:
                self.apply_style(cell, 'ActiveZero')

    def apply_active_total_style(self, ws, row: int, start_col: int, end_col: int):
        """اعمال استایل TOTAL Active Components"""
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            self.apply_style(cell, 'ActiveTotal')